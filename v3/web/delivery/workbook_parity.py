"""Silent old-vs-new workbook parity for scheduled / emailed reports.

Deliveries use the normalized (new) layout. A second workbook is built from the
raw JSON layout, written under the parity folder, compared, and scored in
``view_workbook_parity``. A daily digest email summarizes the scores.
"""

from __future__ import annotations

import io
import logging
import os
import re
import uuid
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from web.data.connection import Database
from web.delivery.layout import apply_layout, expand_clones
from web.reporting.export import build_workbook

log = logging.getLogger(__name__)
EASTERN = ZoneInfo("America/New_York")
_SAFE = re.compile(r"[^A-Za-z0-9._-]+")


@dataclass(frozen=True)
class ParityResult:
    matched: bool
    row_count: int
    old_path: str
    new_path: str
    diff_summary: str


def parity_root(precious_db_path: Path | str) -> Path:
    override = (os.environ.get("VIEW_PARITY_DIR") or "").strip()
    if override:
        return Path(override).expanduser()
    return Path(precious_db_path).expanduser().resolve().parent / "view-parity"


def sheet_grid(xlsx_bytes: bytes) -> tuple[tuple[str, tuple[tuple, ...]], ...]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=False)
    sheets = []
    for name in wb.sheetnames:
        rows = []
        for row in wb[name].iter_rows(values_only=True):
            vals = list(row)
            while vals and vals[-1] is None:
                vals.pop()
            rows.append(tuple(vals))
        sheets.append((name, tuple(rows)))
    wb.close()
    return tuple(sheets)


def diff_sheet_grids(old_sheets, new_sheets) -> list[str]:
    notes: list[str] = []
    old_names = [n for n, _ in old_sheets]
    new_names = [n for n, _ in new_sheets]
    if old_names != new_names:
        return [f"sheet names: old={old_names} new={new_names}"]
    for (name, old_rows), (_, new_rows) in zip(old_sheets, new_sheets):
        if old_rows == new_rows:
            continue
        notes.append(f"{name}: {len(old_rows)} vs {len(new_rows)} rows")
        for i, (a, b) in enumerate(zip(old_rows, new_rows)):
            if a != b:
                notes.append(f"  row {i} old={a!r}")
                notes.append(f"  row {i} new={b!r}")
                if len(notes) > 30:
                    notes.append("  … truncated")
                    return notes
        if len(old_rows) != len(new_rows):
            notes.append(f"  length mismatch ({len(old_rows)} vs {len(new_rows)})")
    return notes


def _safe_name(text: str, fallback: str = "report") -> str:
    s = _SAFE.sub("_", (text or "").strip()).strip("._")
    return (s or fallback)[:80]


def run_parity_files(
    *,
    root: Path,
    report_key: str,
    view_name: str,
    schedule_name: str,
    payload: dict,
    new_layout: dict,
    old_layout: dict,
    delivered_xlsx: bytes | None = None,
) -> ParityResult:
    """Write both workbooks under root/YYYY-MM-DD/ and compare cell grids."""
    now = datetime.now(EASTERN)
    day_dir = root / now.strftime("%Y-%m-%d")
    day_dir.mkdir(parents=True, exist_ok=True)
    stamp = now.strftime("%H%M%S")
    uniq = uuid.uuid4().hex[:8]
    base = (
        f"{stamp}_{uniq}_{_safe_name(report_key)}_"
        f"{_safe_name(view_name or schedule_name or 'view')}"
    )
    new_path = day_dir / f"{base}__new.xlsx"
    old_path = day_dir / f"{base}__old.xlsx"

    if delivered_xlsx is not None:
        new_bytes = delivered_xlsx
    else:
        new_shaped = apply_layout(expand_clones(payload, new_layout), new_layout)
        new_bytes = build_workbook(new_shaped, new_layout)
    old_shaped = apply_layout(expand_clones(payload, old_layout), old_layout)
    old_bytes = build_workbook(old_shaped, old_layout)

    new_path.write_bytes(new_bytes)
    old_path.write_bytes(old_bytes)
    rows = sum(len(t.get("rows") or []) for t in payload.get("tabs") or [])
    notes = diff_sheet_grids(sheet_grid(old_bytes), sheet_grid(new_bytes))
    return ParityResult(
        matched=not notes,
        row_count=rows,
        old_path=str(old_path),
        new_path=str(new_path),
        diff_summary="\n".join(notes)[:4000],
    )


class ViewWorkbookParityRepository:
    def __init__(self, db: Database):
        self.db = db

    def record(
        self,
        *,
        report_key: str,
        view_name: str = "",
        view_id: str | None = None,
        schedule_kind: str = "",
        schedule_id: int | None = None,
        schedule_name: str = "",
        result: ParityResult,
    ) -> int:
        ts = datetime.now(timezone.utc).isoformat()
        with self.db.precious() as conn:
            cur = conn.execute(
                "INSERT INTO view_workbook_parity("
                " created_at, report_key, view_name, view_id, schedule_kind,"
                " schedule_id, schedule_name, matched, row_count, old_path,"
                " new_path, diff_summary)"
                " VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                (
                    ts, report_key, view_name or "", view_id,
                    schedule_kind or "", schedule_id, schedule_name or "",
                    1 if result.matched else 0, result.row_count,
                    result.old_path, result.new_path, result.diff_summary,
                ),
            )
            return int(cur.lastrowid)

    def undigested_before(self, exclusive_end_utc_iso: str) -> list[dict]:
        """All unscored-email rows created before ``exclusive_end_utc_iso``."""
        with self.db.precious() as conn:
            rows = conn.execute(
                "SELECT * FROM view_workbook_parity"
                " WHERE digest_date IS NULL AND created_at < ?"
                " ORDER BY id",
                (exclusive_end_utc_iso,),
            ).fetchall()
            return [dict(r) for r in rows]

    def mark_digested(self, ids: list[int], day: str) -> None:
        if not ids:
            return
        with self.db.precious() as conn:
            conn.executemany(
                "UPDATE view_workbook_parity SET digest_date=? WHERE id=?",
                [(day, i) for i in ids],
            )


def _eastern_day_end_utc_iso(day: str) -> str:
    # Exclusive end = next midnight Eastern.
    from datetime import timedelta

    local = datetime.fromisoformat(f"{day}T00:00:00").replace(tzinfo=EASTERN)
    return (local + timedelta(days=1)).astimezone(timezone.utc).isoformat()


def format_digest(day: str, rows: list[dict]) -> tuple[str, str]:
    matched = sum(1 for r in rows if r.get("matched"))
    diffs = [r for r in rows if not r.get("matched")]
    subject = (
        f"[View parity] {day}: {matched} match, {len(diffs)} diff"
        if rows else f"[View parity] {day}: no dual builds"
    )
    lines = [
        f"Normalized-view workbook parity for {day} (Eastern).",
        f"Total dual builds: {len(rows)}. Match: {matched}. Diff: {len(diffs)}.",
        "",
    ]
    if not rows:
        lines.append("No scheduled/emailed reports recorded a dual build today.")
    for r in rows:
        mark = "MATCH" if r.get("matched") else "DIFF"
        lines.append(
            f"- {mark} {r.get('report_key')} / {r.get('view_name') or '(no view)'} "
            f"schedule={r.get('schedule_name') or '-'} rows={r.get('row_count')}"
        )
        if not r.get("matched") and r.get("diff_summary"):
            for line in str(r["diff_summary"]).splitlines()[:8]:
                lines.append(f"    {line}")
        if r.get("old_path"):
            lines.append(f"    old: {r['old_path']}")
            lines.append(f"    new: {r['new_path']}")
    lines.append("")
    lines.append(
        "Deliveries used the new (normalized) layout. Old JSON workbooks were "
        "written only for this compare."
    )
    return subject, "\n".join(lines)


def raw_legacy_layout(db: Database, legacy_source: str, legacy_id: int) -> dict | None:
    """layout_json from the old table, bypassing live-read hydrate."""
    from web.data.normalized_views import loads_json_object

    with db.precious() as conn:
        row = None
        if legacy_source == "saved_reports":
            row = conn.execute(
                "SELECT layout_json FROM saved_reports WHERE id=?", (legacy_id,),
            ).fetchone()
        elif legacy_source == "company_views":
            row = conn.execute(
                "SELECT layout_json FROM company_views WHERE id=?", (legacy_id,),
            ).fetchone()
        elif legacy_source == "report_defaults":
            vr = conn.execute(
                "SELECT report_key FROM views WHERE legacy_source=? AND legacy_id=?",
                (legacy_source, legacy_id),
            ).fetchone()
            if vr is not None:
                row = conn.execute(
                    "SELECT layout_json FROM report_defaults WHERE report_key=?",
                    (vr["report_key"],),
                ).fetchone()
        if row is None:
            return None
        return loads_json_object(row["layout_json"])
