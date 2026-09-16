"""Column-based views and layout replay (no params_json blobs)."""

from __future__ import annotations

import sqlite3
from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from db import db, init_db
from layout import apply_layout
from main import create_app
from reports import build_payload
from test_home import csrf_headers, login


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("APP_DB_PATH", str(tmp_path / "home.sqlite"))
    monkeypatch.delenv("REPORTING_API_KEY", raising=False)
    monkeypatch.delenv("REPORTING_API_BASE_URL", raising=False)
    with TestClient(create_app()) as test_client:
        yield test_client


def test_views_table_has_no_params_json(client):
    login(client)
    with db() as conn:
        cols = {row[1] for row in conn.execute("PRAGMA table_info(views)").fetchall()}
    assert "params_json" not in cols
    assert "period" in cols
    assert "active_tab_key" in cols


def test_save_view_layout_lives_in_columns(client):
    login(client)
    created = client.post(
        "/api/views",
        json={
            "name": "Hidden name",
            "report_key": "invoiced",
            "kind": "personal",
            "params": {"period": "mtd"},
            "include_period": True,
            "layout": {
                "active": "full_details",
                "order": ["full_details", "invoices"],
                "views": {
                    "full_details": {
                        "hidden": ["CustomerName"],
                        "frozen": ["InvoiceNumber"],
                        "order": ["InvoiceNumber", "CustomerAccount", "Total Invoice"],
                        "sorters": [{"column": "Total Invoice", "dir": "desc"}],
                        "columnFilters": {"CustomerAccount": {"op": "contains", "v": "C-100", "v2": ""}},
                        "group": ["Salesman"],
                        "groups_explicit": True,
                    }
                },
            },
        },
        headers=csrf_headers(client),
    )
    assert created.status_code == 200
    view_id = created.json()["id"]
    with db() as conn:
        tab = conn.execute(
            "SELECT id, groups_explicit FROM layout_tabs WHERE view_id = ? AND tab_key = ?",
            (view_id, "full_details"),
        ).fetchone()
        assert tab is not None
        assert tab["groups_explicit"] == 1
        hidden = [
            row["field"]
            for row in conn.execute(
                "SELECT field FROM layout_columns WHERE tab_id = ? AND hidden = 1",
                (tab["id"],),
            )
        ]
        groups = [
            row["column_name"]
            for row in conn.execute(
                "SELECT column_name FROM layout_tab_groups WHERE tab_id = ? ORDER BY position",
                (tab["id"],),
            )
        ]
        blob = conn.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name='params_json'"
        ).fetchone()
    assert "CustomerName" in hidden
    assert groups == ["Salesman"]
    assert blob is None
    view = client.get("/api/views?report=invoiced").json()["views"]
    saved = next(item for item in view if item["id"] == view_id)
    assert saved["params"]["period"] == "mtd"
    assert saved["layout"]["views"]["full_details"]["hidden"] == ["CustomerName"]
    assert saved["layout"]["views"]["full_details"]["group"] == ["Salesman"]
    html = client.get(f"/reports/invoiced?view={view_id}").text
    assert "data-view-layout" in html
    assert "CustomerName" in html
    assert "groupBySelect" not in html


def test_report_grid_filters_live_in_header_menu(client):
    login(client)
    html = client.get("/reports/invoiced").text
    assert "groupBySelect" not in html
    assert "groupByWrap" not in html
    assert 'id="groupPills"' in html
    js = client.get("/static/js/report-grid.js").text
    assert 'headerFilter: "input"' not in js
    assert "Filter this column" in js
    assert "col-filter-popover" in js
    assert "Group by this column" in js
    assert 'layout: "fitDataTable"' in js
    assert "nestedFieldSeparator: false" in js
    assert "function tableHeight(" in js
    assert 'renderHorizontal: "virtual"' not in js
    assert "if (tableBuilding) return" not in js
    assert "v.frozen.size ? v.frozen.has(field) : idx === 0" not in js


def test_group_string_rejected_on_params_and_layout(client):
    login(client)
    bad_params = client.post(
        "/api/views",
        json={"name": "Bad", "report_key": "invoiced", "params": {"group": "Salesman"}},
        headers=csrf_headers(client),
    )
    assert bad_params.status_code == 400
    bad_layout = client.post(
        "/api/views",
        json={
            "name": "Bad layout",
            "report_key": "invoiced",
            "params": {},
            "layout": {"views": {"full_details": {"group": "Salesman"}}},
        },
        headers=csrf_headers(client),
    )
    assert bad_layout.status_code == 400


def test_xlsx_post_hides_saved_columns(client):
    login(client)
    layout = {
        "views": {
            "full_details": {
                "hidden": ["CustomerName"],
                "order": ["InvoiceNumber", "CustomerAccount", "Total Invoice", "Salesman"],
                "columnFilters": {},
                "group": [],
            }
        }
    }
    res = client.post(
        "/api/reports/invoiced/xlsx",
        json={"layout": layout},
        headers=csrf_headers(client),
    )
    assert res.status_code == 200
    assert res.content[:2] == b"PK"
    book = load_workbook(BytesIO(res.content))
    sheet = book["Full Details"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert "CustomerName" not in headers
    assert "InvoiceNumber" in headers


def test_apply_layout_filters_and_sorts():
    user = {"email": "preview@achimonline.com", "role": "admin", "sales_group": ""}
    payload = build_payload("invoiced", user, {"period": "last_7_days"})
    out = apply_layout(
        payload,
        {
            "views": {
                "full_details": {
                    "hidden": ["SalesmanName"],
                    "sorters": [{"column": "Total Invoice", "dir": "asc"}],
                    "columnFilters": {"CustomerAccount": {"op": "contains", "v": "C-1001", "v2": ""}},
                }
            }
        },
    )
    rows = out["data"]["tabs"]["full_details"]["rows"]
    assert len(rows) == 1
    assert rows[0]["CustomerAccount"] == "C-1001"
    assert "SalesmanName" not in rows[0]


def test_migrate_params_json_into_columns(tmp_path, monkeypatch):
    db_file = tmp_path / "legacy.sqlite"
    monkeypatch.setenv("APP_DB_PATH", str(db_file))
    conn = sqlite3.connect(db_file)
    conn.execute(
        """CREATE TABLE views (
               id INTEGER PRIMARY KEY,
               owner_email TEXT,
               report_key TEXT NOT NULL,
               name TEXT NOT NULL,
               kind TEXT NOT NULL DEFAULT 'personal',
               params_json TEXT,
               include_period INTEGER NOT NULL DEFAULT 0
           )"""
    )
    conn.execute(
        """INSERT INTO views (owner_email, report_key, name, kind, params_json, include_period)
           VALUES ('preview@achimonline.com', 'invoiced', 'Legacy', 'personal',
                   '{"period":"mtd","group":["Salesman"]}', 1)"""
    )
    conn.commit()
    conn.close()
    init_db()
    with db() as conn:
        cols = {row[1] for row in conn.execute("PRAGMA table_info(views)").fetchall()}
        assert "params_json" not in cols
        row = conn.execute("SELECT id, period FROM views WHERE name = 'Legacy'").fetchone()
        assert row["period"] == "mtd"
        groups = conn.execute(
            """SELECT g.column_name FROM layout_tab_groups g
               JOIN layout_tabs t ON t.id = g.tab_id
               WHERE t.view_id = ?""",
            (row["id"],),
        ).fetchall()
    assert [g["column_name"] for g in groups] == ["Salesman"]
