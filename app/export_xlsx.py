"""Build an .xlsx of data.tabs for download."""

from __future__ import annotations

import io

from openpyxl import Workbook

import column_types
import dates

_NUMBER_FORMATS = {
    "money": '"$"#,##0.00_);[Red]\\("$"#,##0.00\\)',
    "int": "#,##0",
    "percent": "0.0%",
    "date": "YYYY-MM-DD",
}


def _as_number(value):
    if isinstance(value, bool) or value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip().replace(",", "").replace("$", "").replace("%", "")
    if not text:
        return None
    try:
        return float(text) if "." in text else int(text)
    except ValueError:
        try:
            return float(text)
        except ValueError:
            return None


def _excel_value(value, col_type: str):
    if value in (None, ""):
        return ""
    if col_type == "date":
        return dates.iso_date(value)
    if col_type in {"money", "int", "percent"}:
        number = _as_number(value)
        if number is None:
            return value
        if col_type == "int":
            return int(number)
        if col_type == "percent" and abs(float(number)) > 1:
            return float(number) / 100.0
        return float(number)
    return value


def workbook_bytes(payload: dict) -> bytes:
    report = payload.get("data") or {}
    tabs = report.get("tabs") or {}
    report_key = str(report.get("report_key") or "")
    book = Workbook()
    is_first_sheet = True
    for key, tab in tabs.items():
        sheet = book.active if is_first_sheet else book.create_sheet()
        is_first_sheet = False
        sheet.title = str(tab.get("name") or key)[:31]
        rows = tab.get("rows") or []
        incoming = tab.get("columns") if isinstance(tab.get("columns"), list) else []
        if not rows and not incoming:
            sheet.append(["(no rows)"])
            continue
        fields = list(rows[0].keys()) if rows else [
            str(col.get("field") or col.get("name") or col) if isinstance(col, dict) else str(col)
            for col in incoming
        ]
        columns = column_types.columns_for(
            fields, incoming=incoming, rows=rows, report_key=report_key
        )
        headers = [col.get("header") or col["field"] for col in columns]
        sheet.append(headers)
        for row in rows:
            values = [_excel_value(row.get(col["field"]), col["type"]) for col in columns]
            sheet.append(values)
            written = sheet[sheet.max_row]
            for cell, col in zip(written, columns):
                fmt = _NUMBER_FORMATS.get(col["type"])
                if fmt and isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = fmt
                elif col["type"] == "date" and cell.value not in (None, ""):
                    cell.number_format = _NUMBER_FORMATS["date"]
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()
